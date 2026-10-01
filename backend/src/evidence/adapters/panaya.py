"""Panaya ETL adapter (ADR-017), rewritten against a real ~2.8GB customer export.

No reviewed Panaya sample existed when this adapter was first written (ADR-017's original
`<PanayaExport>` profile was an invented placeholder). Validated later against an actual export,
the real structure is a single `<ROOT_ELEMENT>` containing `<HEADER .../>` (system metadata as
attributes), `<ETL_RUN_PARAMS>`, and ~90 sibling **sections** (`<REPOSITORY_OBJECTS>`,
`<PROGRAMS>`, `<WHERE_USED_TABLE>`, `<SCI_HANA_ISSUES>`, ...), each holding thousands to millions
of flat, attribute-only child records. The file is **not** wrapped in a ZIP in practice — a bare
multi-gigabyte `.xml` — though a ZIP-wrapped variant is still supported in case some export
configuration produces one.

Two real, load-bearing quirks this adapter handles that the invented profile never anticipated:

1. **Not strictly well-formed XML.** Raw ABAP/text content can carry literal illegal control
   bytes (confirmed: a raw `0x1C` inside a `<![CDATA[...]]>`-adjacent comment) *and* escaped
   illegal numeric character references (confirmed: `&#0;` for a NUL byte in a garbled field).
   Both are outright forbidden by the XML 1.0 `Char` production and abort `expat` with no way to
   resume. `evidence.adapters.base.SanitizingXMLStream` strips both, streaming, before the parser
   ever sees them — dropping only the illegal placeholder markers, never fabricating content.
2. **Scale.** ~90 sections, several in the single-digit millions of rows (confirmed:
   `SCI_HANA_PERFORMANCE_ISSUES_DETAILS` ~5M, `WHERE_USED_TABLE` ~1.2M). Importing every row of
   every section is not a PoC-appropriate goal. This adapter extracts a small **curated** subset
   of sections with verified, correlatable attribute shapes (`_CURATED_SECTIONS`); every section
   — curated or not — is still counted in `inspect()`'s manifest so nothing is invisible, and a
   curated section's rows are capped (`_MAX_RECORDS_PER_SECTION`) with an explicit truncation
   warning rather than an unbounded import. Extending the curated set to another section is a
   one-line dict entry once its real attribute shape has been verified — never guessed.

Because random access into a multi-gigabyte XML isn't practical, each batch = one curated
section, re-streamed from the start of the file and stopped as soon as that section's closing
tag is consumed (sections earlier in the document finish fast; a late section like
`WHERE_USED_TABLE` costs close to a full pass — acceptable for a one-time background import).
"""
from __future__ import annotations

import datetime
import hashlib
import zipfile
from contextlib import contextmanager
from pathlib import Path

import openpyxl
from defusedxml import ElementTree as ET
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from evidence.adapters.base import ImportBatchOutcome, ImportBatchPlan, InspectionResult, SanitizingXMLStream
from persistence.models import EvidenceDataset, EvidenceRecord

DATASET_TYPE = "PANAYA_ETL"
IMPORTER_VERSION = "2.0"

_SIGNATURE = b"EXPORT_TOOL_VERSION"
_MAX_RECORDS_PER_SECTION = 20000

# A second, real, independently-verified Panaya export format (ADR-017 amendment, SPRINT-20): a
# flat XLSX "usage/repository" report — one sheet, one header row, one row per SAP object, with a
# per-object usage level (Unused/Normally Used/...) and origin (Customer/SAP Standard/
# 3rdPartDomain) that no curated XML section above carries. Verified against a real
# ~143k-row customer export (`export T-systems.xlsx`). Detected by header-row *signature*
# (tolerant of column reordering/extra columns), never by filename or an exact-match requirement
# — only one real sample has been reviewed, and ADR-016/ADR-017 both forbid hard-coding the one
# reviewed sample as the only valid schema. Kept under this same `PANAYA_ETL` dataset_type/adapter
# module rather than a second Panaya dataset_type, since it is still Panaya-sourced landscape data
# about the same kind of SAP objects.
_USAGE_XLSX_HEADER_SIGNATURE = {"OBJECT NAME", "OBJECT TYPE", "USAGE LEVEL", "ORIGIN"}
_USAGE_XLSX_FORMAT = "xlsx_usage_report"
_USAGE_XLSX_BATCH_SIZE = 20000

# Verified against the real export — only sections with a confirmed, correlatable attribute
# shape are listed. Everything else is counted (inspect()'s manifest) but not persisted as
# records yet; add an entry here once a new section's real shape is verified.
_CURATED_SECTIONS: dict[str, dict[str, str]] = {
    "REPOSITORY_OBJECTS": {
        "capability": "TECHNICAL_OBJECT_METADATA", "record_type": "panaya_repository_object",
        "object_name_attr": "OBJ_NAME", "object_type_attr": "OBJECT", "package_attr": "DEVCLASS",
    },
    "PROGRAMS": {
        "capability": "TECHNICAL_OBJECT_METADATA", "record_type": "panaya_program",
        "object_name_attr": "NAME",
    },
    "FUNCTIONS": {
        "capability": "TECHNICAL_OBJECT_METADATA", "record_type": "panaya_function",
        "object_name_attr": "FUNCNAME",
    },
    "MODIFICATIONS": {
        "capability": "SOURCE_CODE", "record_type": "panaya_modification",
        "object_name_attr": "OBJ_NAME", "object_type_attr": "OBJ_TYPE",
    },
    "NOTES_HEADER": {
        "capability": "S4_CONVERSION_SIGNAL", "record_type": "panaya_sap_note",
    },
    # Verified 2026-09-30 against a real ~2.8GB customer export (SPRINT-19): TS_HANAA_SCI_CHECK
    # really exports OBJTYPE/OBJNAME/DEVCLASS (same shape as REPOSITORY_OBJECTS) plus a DETAILS_REF
    # key, preserved in normalized_payload alongside every other attribute — no special join code
    # needed for a details row to be traced back to its parent check.
    "SCI_HANA_ISSUES": {
        "capability": "S4_CONVERSION_SIGNAL", "record_type": "panaya_sci_hana_issue",
        "object_name_attr": "OBJNAME", "object_type_attr": "OBJTYPE", "package_attr": "DEVCLASS",
    },
    # TS_HANA_SCI_CHECK_DETAILS really exports DETAILS_REF/TEXT/INCLUDE/LINE (no PROGRAM attribute
    # observed in the real export) — INCLUDE is the only correlatable object name here.
    "SCI_HANA_ISSUES_DETAILS": {
        "capability": "S4_CONVERSION_SIGNAL", "record_type": "panaya_sci_hana_issue_detail",
        "object_name_attr": "INCLUDE",
    },
    # WBCROSSGT/CROSS (SELECT * shape) really export OTYPE/NAME/INCLUDE/DIRECT. INCLUDE (the
    # referencing program/include) is a clean, correlatable object name — NAME is deliberately NOT
    # used here: it is a compound cross-reference token whose structure depends on OTYPE (e.g. a
    # data-object/variable path like "CLASS\ME:METHOD\DA:VAR", not a bare object name), and treating
    # it as one would manufacture false-confidence correlations (the exact failure ADR-017 forbids).
    "WHERE_USED_TABLE": {
        "capability": "DEPENDENCY_SIGNAL", "record_type": "panaya_where_used",
        "object_name_attr": "INCLUDE",
    },
}


def detect(filename: str, file_path: Path) -> bool:
    lower = filename.lower()
    # .xlsx is only ever matched by real header-row content (never by filename) — the XML/ZIP
    # fast path below is deliberately not reused here, since a file merely named "...panaya...xlsx"
    # is not necessarily this real usage-report shape.
    if lower.endswith(".xlsx"):
        return _usage_xlsx_header(file_path) is not None
    if "panaya" in lower:
        return True
    if not (lower.endswith(".xml") or lower.endswith(".zip")):
        return False
    try:
        with _open_xml_stream(file_path) as stream:
            head = stream.read(65536)
    except (zipfile.BadZipFile, ValueError, OSError):
        return False
    return _SIGNATURE in head


def _usage_xlsx_header(file_path: Path) -> list[str] | None:
    """Return the header row (original casing/order, blanks as `""`) if `file_path` is a
    workbook whose first sheet's header row carries the known usage-report column signature —
    tolerant of extra/reordered columns, never an exact-match requirement (ADR-016/ADR-017's
    shared "never hard-code the one reviewed sample as the only valid schema" principle). `None`
    if the file can't be read as a workbook or the signature isn't present."""
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        return None
    try:
        ws = wb.worksheets[0]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    finally:
        wb.close()
    if not header_row:
        return None
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    if _USAGE_XLSX_HEADER_SIGNATURE.issubset({h.upper() for h in headers if h}):
        return headers
    return None


def _usage_xlsx_cell(value: object) -> object:
    """JSONB-safe cell value — openpyxl can hand back `datetime.date`/`datetime.datetime` for a
    date-formatted cell; every value observed in the real sample was already a plain string, but
    normalized_payload must stay JSON-serializable regardless of a future export's cell
    formatting."""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return value


def _local_tag(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _find_primary_xml_member(zf: zipfile.ZipFile) -> str | None:
    xml_members = [n for n in zf.namelist() if n.lower().endswith(".xml")]
    if not xml_members:
        return None
    return max(xml_members, key=lambda n: zf.getinfo(n).file_size)


@contextmanager
def _open_xml_stream(file_path: Path):
    """Yield a `SanitizingXMLStream` over the primary XML, whether `file_path` is a bare
    `.xml` (the real-world case) or a ZIP wrapping one (kept for defense in depth)."""
    if file_path.suffix.lower() == ".zip":
        with zipfile.ZipFile(file_path) as zf:
            xml_name = _find_primary_xml_member(zf)
            if xml_name is None:
                raise ValueError("No XML member found in ZIP")
            try:
                fh = zf.open(xml_name)
            except NotImplementedError as exc:
                method = zipfile.compressor_names.get(zf.getinfo(xml_name).compress_type, "unknown")
                raise NotImplementedError(
                    f"ZIP member {xml_name!r} uses compression method '{method}', which Python's "
                    f"standard zipfile module cannot decode ({exc}). Re-export as a bare .xml (no "
                    f"ZIP) or with a supported method (store/deflate/bzip2/lzma), or upload the "
                    f"already-unzipped .xml directly if you have it."
                ) from exc
            with fh:
                yield SanitizingXMLStream(fh)
    else:
        with open(file_path, "rb") as raw:
            yield SanitizingXMLStream(raw)


def inspect(file_path: Path) -> InspectionResult:
    if file_path.suffix.lower() == ".xlsx":
        return _inspect_usage_xlsx(file_path)

    header_attrs: dict[str, str] = {}
    section_counts: dict[str, int] = {}
    depth = 0
    current_section: str | None = None

    try:
        with _open_xml_stream(file_path) as stream:
            for event, elem in ET.iterparse(stream, events=("start", "end")):
                tag = _local_tag(elem.tag)
                if event == "start":
                    depth += 1
                    if depth == 2:
                        current_section = tag
                        section_counts.setdefault(tag, 0)
                        if tag == "HEADER":
                            header_attrs = dict(elem.attrib)
                else:
                    if depth == 3 and current_section is not None:
                        section_counts[current_section] += 1
                    if depth >= 2:
                        elem.clear()
                    depth -= 1
    except (zipfile.BadZipFile, ValueError) as exc:
        return InspectionResult(dataset_type=DATASET_TYPE, display_name=file_path.name, warnings=[str(exc)])
    except (ET.ParseError, DefusedXmlException) as exc:
        return InspectionResult(
            dataset_type=DATASET_TYPE, display_name=file_path.name,
            warnings=[f"XML parse error even after illegal-character sanitization: {exc}"],
        )
    except (NotImplementedError, RuntimeError) as exc:
        # zipfile.open() raises these for an unsupported compression method or an
        # encrypted entry (only reachable via the ZIP-wrapped fallback path) —
        # `_open_xml_stream` already names the specific method for NotImplementedError.
        return InspectionResult(dataset_type=DATASET_TYPE, display_name=file_path.name, warnings=[str(exc)])
    except OSError as exc:
        return InspectionResult(dataset_type=DATASET_TYPE, display_name=file_path.name, warnings=[f"Cannot read file: {exc}"])

    capabilities = sorted(
        {info["capability"] for name, info in _CURATED_SECTIONS.items() if section_counts.get(name)}
    )
    warnings: list[str] = []
    if not section_counts:
        warnings.append("No recognized Panaya ETL sections found")

    return InspectionResult(
        dataset_type=DATASET_TYPE,
        display_name=file_path.name,
        capabilities=capabilities,
        manifest={
            "header": header_attrs,
            "section_counts": section_counts,
            "curated_sections": {n: section_counts.get(n, 0) for n in _CURATED_SECTIONS},
        },
        source_system_hint=header_attrs.get("SYSTEM_ID"),
        source_client_hint=header_attrs.get("CLIENT"),
        warnings=warnings,
    )


def _inspect_usage_xlsx(file_path: Path) -> InspectionResult:
    headers = _usage_xlsx_header(file_path)
    if headers is None:
        return InspectionResult(
            dataset_type=DATASET_TYPE, display_name=file_path.name,
            warnings=["XLSX header row does not carry the recognized Panaya usage-report column signature"],
        )
    col_index = {h.upper(): i for i, h in enumerate(headers) if h}

    def _counter_for(column: str) -> dict[str, int] | None:
        return {} if column in col_index else None

    origin_counts = _counter_for("ORIGIN")
    usage_counts = _counter_for("USAGE LEVEL")
    type_counts = _counter_for("OBJECT TYPE")
    module_counts = _counter_for("MODULE")

    row_count = 0
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        try:
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_count += 1
                for counts, column in (
                    (origin_counts, "ORIGIN"), (usage_counts, "USAGE LEVEL"),
                    (type_counts, "OBJECT TYPE"), (module_counts, "MODULE"),
                ):
                    if counts is None:
                        continue
                    idx = col_index[column]
                    value = str(row[idx]) if idx < len(row) and row[idx] not in (None, "") else "(vazio)"
                    counts[value] = counts.get(value, 0) + 1
        finally:
            wb.close()
    except (InvalidFileException, zipfile.BadZipFile, OSError) as exc:
        return InspectionResult(dataset_type=DATASET_TYPE, display_name=file_path.name, warnings=[f"Cannot read XLSX: {exc}"])

    return InspectionResult(
        dataset_type=DATASET_TYPE,
        display_name=file_path.name,
        capabilities=["TECHNICAL_OBJECT_METADATA", "USAGE_SIGNAL"],
        manifest={
            "source_format": _USAGE_XLSX_FORMAT,
            "headers": headers,
            "row_count": row_count,
            "origin_counts": origin_counts or {},
            "usage_level_counts": usage_counts or {},
            "object_type_counts": type_counts or {},
            "module_counts": module_counts or {},
        },
        # No system/client header metadata exists in this format (unlike the XML export's
        # <HEADER SYSTEM_ID=... CLIENT=... />) — a real limitation of the format, not guessed.
        source_system_hint=None,
        source_client_hint=None,
        warnings=[],
    )


def plan_batches(file_path: Path, inspection: InspectionResult) -> list[ImportBatchPlan]:
    if inspection.manifest.get("source_format") == _USAGE_XLSX_FORMAT:
        row_count = inspection.manifest.get("row_count", 0)
        return [
            ImportBatchPlan(
                batch_key=f"usage_rows:{start}",
                payload={"start_row": start, "end_row": min(start + _USAGE_XLSX_BATCH_SIZE, row_count)},
            )
            for start in range(0, row_count, _USAGE_XLSX_BATCH_SIZE)
        ]
    counts = inspection.manifest.get("section_counts", {})
    return [
        ImportBatchPlan(batch_key=name, payload={"section": name})
        for name in _CURATED_SECTIONS
        if counts.get(name, 0) > 0
    ]


def _import_usage_xlsx_batch(
    file_path: Path, dataset: EvidenceDataset, batch: ImportBatchPlan, session: Session
) -> ImportBatchOutcome:
    start_row = batch.payload["start_row"]
    end_row = batch.payload["end_row"]

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in header_row]

        imported = 0
        idx = start_row
        # Sheet rows are 1-indexed with row 1 as the header; `start_row`/`end_row` are 0-indexed
        # data-row offsets, so the first data row (offset 0) is sheet row 2.
        for row in ws.iter_rows(min_row=2 + start_row, max_row=1 + end_row, values_only=True):
            attrs = {
                headers[i]: _usage_xlsx_cell(row[i])
                for i in range(len(headers))
                if i < len(row) and row[i] not in (None, "")
            }
            source_key = f"usage_row:{idx}"
            fingerprint = hashlib.sha256(source_key.encode()).hexdigest()[:32]
            already = session.execute(
                select(EvidenceRecord.id).where(
                    EvidenceRecord.dataset_id == dataset.id, EvidenceRecord.record_fingerprint == fingerprint
                )
            ).first()
            if already is None:
                session.add(
                    EvidenceRecord(
                        dataset_id=dataset.id,
                        record_type="panaya_usage_object",
                        capability="USAGE_SIGNAL",
                        source_key=source_key,
                        record_fingerprint=fingerprint,
                        object_name=attrs.get("OBJECT NAME") or None,
                        object_type=attrs.get("OBJECT TYPE") or None,
                        package_name=attrs.get("PACKAGE") or None,
                        normalized_payload=attrs,
                        source_locator={"row_index": idx},
                    )
                )
                imported += 1
                if imported % 500 == 0:
                    session.flush()
            idx += 1
    finally:
        wb.close()

    session.flush()
    return ImportBatchOutcome(records_imported=imported, warnings=[])


def import_batch(
    file_path: Path, dataset: EvidenceDataset, batch: ImportBatchPlan, session: Session
) -> ImportBatchOutcome:
    if "start_row" in batch.payload:
        return _import_usage_xlsx_batch(file_path, dataset, batch, session)

    section = batch.payload["section"]
    info = _CURATED_SECTIONS[section]
    rows: list[tuple[int, dict]] = []
    depth = 0
    current_section: str | None = None
    row_index = 0
    truncated = False

    with _open_xml_stream(file_path) as stream:
        for event, elem in ET.iterparse(stream, events=("start", "end")):
            tag = _local_tag(elem.tag)
            if event == "start":
                depth += 1
                if depth == 2:
                    current_section = tag
            else:
                if depth == 3 and current_section == section:
                    if row_index < _MAX_RECORDS_PER_SECTION:
                        rows.append((row_index, dict(elem.attrib)))
                    else:
                        truncated = True
                    row_index += 1
                if depth >= 2:
                    elem.clear()
                depth -= 1
                if depth == 1 and current_section == section:
                    break  # done with the target section — no need to read the rest of the file

    imported = 0
    for idx, attrs in rows:
        source_key = f"{section}:{idx}"
        fingerprint = hashlib.sha256(source_key.encode()).hexdigest()[:32]
        already = session.execute(
            select(EvidenceRecord.id).where(
                EvidenceRecord.dataset_id == dataset.id, EvidenceRecord.record_fingerprint == fingerprint
            )
        ).first()
        if already is not None:
            continue

        object_name = attrs.get(info["object_name_attr"]) if "object_name_attr" in info else None
        object_type = attrs.get(info["object_type_attr"]) if "object_type_attr" in info else None
        package_name = attrs.get(info["package_attr"]) if "package_attr" in info else None

        session.add(
            EvidenceRecord(
                dataset_id=dataset.id,
                record_type=info["record_type"],
                capability=info["capability"],
                source_key=source_key,
                record_fingerprint=fingerprint,
                object_name=object_name or None,
                object_type=object_type or None,
                package_name=package_name or None,
                normalized_payload=attrs,
                source_locator={"section": section, "row_index": idx},
            )
        )
        imported += 1
        if imported % 500 == 0:
            session.flush()

    session.flush()
    warnings = (
        [f"Section {section} truncated to {_MAX_RECORDS_PER_SECTION} records — more rows exist in the source"]
        if truncated else []
    )
    return ImportBatchOutcome(records_imported=imported, warnings=warnings)
