"""SPRINT-08 (post-completion extension) — SAP Readiness Check (.docx) adapter."""
from __future__ import annotations

from pathlib import Path

import docx
from sqlalchemy import select, text

from evidence.adapters import readiness_check
from persistence.database import get_session_factory
from persistence.models import (
    Assessment,
    AssessmentStatus,
    Client,
    EvidenceDataset,
    EvidenceDatasetStatus,
    EvidenceRecord,
)


def _make_docx(tmp_path: Path, name: str = "readiness_check.docx") -> Path:
    doc = docx.Document()
    doc.add_paragraph("SAP Readiness Check for SAP S/4HANA Conversion")

    meta = doc.add_table(rows=2, cols=2)
    meta.rows[0].cells[0].text, meta.rows[0].cells[1].text = "SID", "Database System"
    meta.rows[1].cells[0].text, meta.rows[1].cells[1].text = "QAS", "MSSQL"

    simp = doc.add_table(rows=3, cols=5)
    hdr = ["LoB–BusinessArea", "Title", "Effort Ranking", "SAP Note", "Relevance"]
    for i, h in enumerate(hdr):
        simp.rows[0].cells[i].text = h
    simp.rows[1].cells[0].text, simp.rows[1].cells[1].text = "Finance", "New Asset Accounting"
    simp.rows[1].cells[2].text, simp.rows[1].cells[3].text = "M", "2270388"
    simp.rows[1].cells[4].text = "Relevant"
    simp.rows[2].cells[0].text, simp.rows[2].cells[1].text = "Logistics", "Material Number Length"
    simp.rows[2].cells[2].text, simp.rows[2].cells[3].text = "L", "2267140"
    simp.rows[2].cells[4].text = "Relevant"

    code = doc.add_table(rows=2, cols=3)
    code.rows[0].cells[0].text = "Check Type"
    code.rows[0].cells[1].text = "Number Issues Found"
    code.rows[0].cells[2].text = "Description"
    code.rows[1].cells[0].text, code.rows[1].cells[1].text = "Syntax Error", "12"
    code.rows[1].cells[2].text = "Custom code syntax errors"

    vol = doc.add_table(rows=3, cols=3)
    vol.rows[0].cells[0].text = "Table"
    vol.rows[0].cells[1].text = "Column Loadable Size in GiB"
    vol.rows[0].cells[2].text = "Estimated Number of Records"
    vol.rows[1].cells[0].text, vol.rows[1].cells[1].text, vol.rows[1].cells[2].text = "DBTABLOG", "230,0", "651620305"
    vol.rows[2].cells[0].text, vol.rows[2].cells[1].text, vol.rows[2].cells[2].text = "ACDOCA", "77,6", "1699689986"

    unrelated = doc.add_table(rows=2, cols=2)
    unrelated.rows[0].cells[0].text, unrelated.rows[0].cells[1].text = "Topic", "Link"
    unrelated.rows[1].cells[0].text, unrelated.rows[1].cells[1].text = "Fiori Apps", "https://example.com"

    path = tmp_path / name
    doc.save(path)
    return path


def _make_assessment(session) -> int:
    c = Client(name="Readiness Check Adapter Test Client")
    session.add(c)
    session.flush()
    a = Assessment(client_id=c.id, name="Readiness Check Adapter Test", status=AssessmentStatus.CREATED.value)
    session.add(a)
    session.commit()
    return a.id


def _make_dataset(session, assessment_id: int) -> int:
    ds = EvidenceDataset(
        assessment_id=assessment_id, dataset_type=readiness_check.DATASET_TYPE,
        display_name="readiness_check.docx", source_filename="readiness_check.docx",
        source_sha256="b" * 64, source_size_bytes=10, importer_name="readiness_check",
        importer_version=readiness_check.IMPORTER_VERSION, status=EvidenceDatasetStatus.IMPORTING.value,
    )
    session.add(ds)
    session.flush()
    return ds.id


def _cleanup(session, assessment_id: int) -> None:
    session.execute(
        text("DELETE FROM client WHERE id = (SELECT client_id FROM assessment WHERE id = :aid)"),
        {"aid": assessment_id},
    )
    session.commit()


def test_detect_matches_filename_or_content(tmp_path: Path) -> None:
    path = _make_docx(tmp_path)
    assert readiness_check.detect("SAP Readiness Check for SAP S4HANA Conversion.docx", path)
    assert readiness_check.detect("export.docx", path)  # content sniffing (title paragraph)
    assert not readiness_check.detect("export.txt", path)

    plain = docx.Document()
    plain.add_paragraph("Just a normal document, nothing SAP-related")
    other_path = tmp_path / "other.docx"
    plain.save(other_path)
    assert not readiness_check.detect("other.docx", other_path)


def test_inspect_classifies_tables_by_header(tmp_path: Path) -> None:
    path = _make_docx(tmp_path)
    result = readiness_check.inspect(path)

    assert result.capabilities == ["S4_CONVERSION_SIGNAL", "USAGE_SIGNAL"]  # sorted, deduplicated
    assert result.manifest["table_count"] == 5
    assert result.manifest["recognized_table_counts"] == {
        "system_metadata": 1, "simplification_item": 1, "custom_code_check": 1, "table_volume": 1,
    }
    assert result.manifest["unrecognized_tables"] == 1
    assert result.source_system_hint == "QAS"


def test_import_persists_classified_rows_with_correlatable_object_names(tmp_path: Path) -> None:
    path = _make_docx(tmp_path)
    inspection = readiness_check.inspect(path)
    batches = readiness_check.plan_batches(path, inspection)
    assert len(batches) == 4  # system_metadata + simplification_item + custom_code_check + table_volume

    with get_session_factory()() as session:
        try:
            asmnt_id = _make_assessment(session)
            ds_id = _make_dataset(session, asmnt_id)
            dataset = session.get(EvidenceDataset, ds_id)

            total = 0
            for batch in batches:
                outcome = readiness_check.import_batch(path, dataset, batch, session)
                total += outcome.records_imported
            session.commit()
            assert total == 5  # 2 simplification items + 1 custom code check + 2 table volumes

            records = list(session.scalars(select(EvidenceRecord).where(EvidenceRecord.dataset_id == ds_id)))
            by_type: dict[str, int] = {}
            for r in records:
                by_type[r.record_type] = by_type.get(r.record_type, 0) + 1
            assert by_type["readiness_check_simplification_item"] == 2
            assert by_type["readiness_check_custom_code_check"] == 1
            assert by_type["readiness_check_table_volume"] == 2

            table_records = [r for r in records if r.record_type == "readiness_check_table_volume"]
            assert {r.object_name for r in table_records} == {"DBTABLOG", "ACDOCA"}
            simp_records = [r for r in records if r.record_type == "readiness_check_simplification_item"]
            assert all(r.object_name is None for r in simp_records)
            assert any(r.normalized_payload.get("Title") == "New Asset Accounting" for r in simp_records)

            # Idempotent re-run of a batch that actually persisted records.
            simplification_batch = next(b for b in batches if b.payload["kind"] == "simplification_item")
            outcome_again = readiness_check.import_batch(path, dataset, simplification_batch, session)
            session.commit()
            assert outcome_again.records_imported == 0
        finally:
            _cleanup(session, asmnt_id)
